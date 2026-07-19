# -*- coding: utf-8 -*-
"""Compara os valores calculados pelo Excel (cache) com o esperado (zscore_ref)."""
import os, json
from openpyxl import load_workbook

DIR = os.path.dirname(os.path.abspath(__file__))
TMP = os.path.join(DIR, "_verif_temp.xlsx")
EXP = os.path.join(DIR, "_verif_expected.json")

exp = json.load(open(EXP))
wb = load_workbook(TMP, data_only=True)
ws = wb["Calculo"]

# colunas: K=11 zIMC,L=12 dIMC,M=13 zEst,N=14 dEst,O=15 zPeso,P=16 dPeso,Q=17 zPE,R=18 dPE
COLZ = {"z_imc": 11, "z_est": 13, "z_peso": 15, "z_pe": 17}
COLD = {"d_imc": 12, "d_est": 14, "d_peso": 16, "d_pe": 18}


def norm_z(v):
    if v is None or v == "" or v == "—":
        return None
    return round(float(v), 2)


erros = 0
print(f"{'ID':6} {'mes':>4} | {'idx':9} {'excel':>8} {'ref':>8}  ok")
for caso in exp["casos"]:
    r = caso["row"]
    for key, col in COLZ.items():
        excel = ws.cell(row=r, column=col).value
        # detectar erros de formula
        if isinstance(excel, str) and excel.startswith("#"):
            print(f"  ERRO FORMULA {caso['ID']} {key}: {excel}")
            erros += 1
            continue
        ez = norm_z(excel)
        rz = caso[key]
        ok = (ez is None and rz is None) or (ez is not None and rz is not None and abs(ez - rz) <= 0.011)
        if not ok:
            erros += 1
            print(f"{caso['ID']:6} {str(caso['meses']):>4} | {key:9} {str(ez):>8} {str(rz):>8}  {'OK' if ok else 'XXX'}")
    # diagnosticos
    for key, col in COLD.items():
        excel = ws.cell(row=r, column=col).value or ""
        ref = caso[key] or ""
        if str(excel).strip() != str(ref).strip():
            erros += 1
            print(f"  DIAG DIFF {caso['ID']} {key}: excel='{excel}' ref='{ref}'")

# Consulta
cs = wb["Consulta"]
# resultados em G4..G8 (IMC,zIMC,zEst,zPeso,zPE) -> linhas 4..8 col7
con_excel = {"z_imc": cs.cell(row=5, column=7).value, "z_est": cs.cell(row=6, column=7).value,
             "z_peso": cs.cell(row=7, column=7).value, "z_pe": cs.cell(row=8, column=7).value}
print("\nConsulta (caso obeso):")
for k in ["z_imc", "z_est", "z_peso", "z_pe"]:
    ez = norm_z(con_excel[k]); rz = exp["consulta"][k]
    ok = (ez is None and rz is None) or (ez is not None and rz is not None and abs(ez - rz) <= 0.011)
    if not ok: erros += 1
    print(f"  {k:7} excel={ez} ref={rz} {'OK' if ok else 'XXX'}")

# Ficha (C001 tem 3 medidas) -> contar linhas preenchidas na coluna Data (col3) a partir da linha 9
fc = wb["Ficha"]
nhist = sum(1 for r in range(9, 39) if fc.cell(row=r, column=3).value not in (None, ""))
print(f"\nFicha C001: {nhist} medidas no historico (esperado 3)  {'OK' if nhist == 3 else 'XXX'}")
if nhist != 3: erros += 1

# Painel
pn = wb["Painel"]
tot = pn.cell(row=4, column=5).value
print(f"Painel total avaliados: {tot} (esperado 16)")

print(f"\n{'== TUDO OK ==' if erros == 0 else f'== {erros} DIVERGENCIAS =='}")
