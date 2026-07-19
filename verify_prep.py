# -*- coding: utf-8 -*-
"""Prepara arquivo temporario com casos-teste injetados na aba Calculo +
expected.json com os valores esperados (do zscore_ref validado)."""
import os, json, datetime as dt
from openpyxl import load_workbook
import zscore_ref as zr

DIR = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(DIR, "Calculadora_OMS.xlsx")
TMP = os.path.join(DIR, "_verif_temp.xlsx")
EXP = os.path.join(DIR, "_verif_expected.json")

D = dt.date
# (ID, Nome, SexoTxt, DN, Data, Peso, Altura, Medida)
CASOS = [
    ("C001", "Ana",   "Feminino", D(2021, 6, 15), D(2024, 6, 15), 14.0, 96.0, "Em pe"),   # 36m
    ("C001", "Ana",   "Feminino", D(2021, 6, 15), D(2023, 6, 15), 11.5, 86.0, "Em pe"),   # 24m (mesma crianca)
    ("C001", "Ana",   "Feminino", D(2021, 6, 15), D(2022, 6, 15), 9.0, 74.0, "Deitado"),  # 12m
    ("N000", "RN m",  "Masculino", D(2024, 6, 15), D(2024, 6, 15), 3.3, 49.9, "Deitado"),  # 0d
    ("B003", "Bebe",  "Masculino", D(2024, 3, 15), D(2024, 6, 15), 6.4, 61.4, "Deitado"),  # ~3m
    ("M024", "Boy24", "Masculino", D(2022, 6, 15), D(2024, 6, 15), 12.2, 87.1, "Em pe"),   # 24m
    ("G060", "Girl5", "Feminino", D(2019, 6, 15), D(2024, 6, 15), 18.0, 109.0, "Em pe"),   # 60m
    ("B061", "Boy5a1","Masculino", D(2019, 5, 15), D(2024, 6, 15), 18.3, 110.0, "Em pe"),  # 61m
    ("G120", "Girl10","Feminino", D(2014, 6, 15), D(2024, 6, 15), 32.0, 138.0, "Em pe"),   # 120m
    ("B132", "Boy11", "Masculino", D(2013, 6, 15), D(2024, 6, 15), 38.0, 145.0, "Em pe"),  # 132m (peso/idade -> —)
    ("G228", "Girl19","Feminino", D(2005, 6, 15), D(2024, 6, 15), 55.0, 163.0, "Em pe"),   # 228m
    ("B229", "Boy19a1","Masculino", D(2005, 5, 15), D(2024, 6, 15), 60.0, 170.0, "Em pe"), # 229m (-> —)
    ("OBE8", "Obeso", "Masculino", D(2016, 6, 15), D(2024, 6, 15), 45.0, 125.0, "Em pe"),  # 96m IMC z>3
    ("WAS3", "Magra", "Feminino", D(2021, 6, 15), D(2024, 6, 15), 8.0, 100.0, "Em pe"),    # 36m peso/est z<-3
    ("LIE2", "Deit2", "Masculino", D(2022, 6, 15), D(2024, 6, 15), 13.0, 88.0, "Deitado"), # 24m medido deitado -> -0.7
    ("STD1", "EmPe1", "Feminino", D(2023, 6, 15), D(2024, 6, 15), 9.5, 73.0, "Em pe"),     # 12m medido em pe -> +0.7
]

SEXMAP = {"Masculino": 1, "Feminino": 2}

wb = load_workbook(SRC)
ws = wb["Calculo"]
expected = []
for i, c in enumerate(CASOS):
    r = 2 + i
    ID, Nome, Sx, DN, Data, P, A, Med = c
    ws.cell(row=r, column=1, value=ID)
    ws.cell(row=r, column=2, value=Nome)
    ws.cell(row=r, column=3, value=Sx)
    ws.cell(row=r, column=4, value=DN)
    ws.cell(row=r, column=5, value=Data)
    ws.cell(row=r, column=6, value=P)
    ws.cell(row=r, column=7, value=A)
    ws.cell(row=r, column=8, value=Med)
    e = zr.calc(SEXMAP[Sx], DN, Data, P, A, Med)
    expected.append({
        "row": r, "ID": ID, "meses": e["meses"],
        "z_imc": e["z_imc"], "z_est": e["z_est"], "z_peso": e["z_peso"], "z_pe": e["z_pe"],
        "d_imc": zr.classif_imc(e["z_imc"], e["meses"]) if e["meses"] is not None else "",
        "d_est": zr.classif_est(e["z_est"]),
        "d_peso": zr.classif_peso(e["z_peso"]),
        "d_pe": zr.classif_pe(e["z_pe"]),
    })

# Ficha: seleciona C001 (3 medidas) ; Consulta: caso obeso
wb["Ficha"]["C4"] = "C001"
cs = wb["Consulta"]
cs["C4"] = "Masculino"; cs["C5"] = D(2016, 6, 15); cs["C6"] = D(2024, 6, 15)
cs["C7"] = 45.0; cs["C8"] = 125.0; cs["C9"] = "Em pe"
exp_consulta = zr.calc(1, D(2016, 6, 15), D(2024, 6, 15), 45.0, 125.0, "Em pe")

wb.save(TMP)
json.dump({"casos": expected,
           "consulta": {"z_imc": exp_consulta["z_imc"], "z_est": exp_consulta["z_est"],
                        "z_peso": exp_consulta["z_peso"], "z_pe": exp_consulta["z_pe"]},
           "n": len(CASOS)},
          open(EXP, "w"))
print(f"Temp salvo: {TMP}  ({len(CASOS)} casos)")
